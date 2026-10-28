"""Optional, fixed-layout intermediate dilutions for the Screen workflow.

This module plans preparation and output artifacts. It never randomizes the
intermediate plate or modifies the input design workbook.
"""
from __future__ import annotations

from collections import OrderedDict
from dataclasses import asdict, dataclass
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
import json
import math
from pathlib import Path
import re
from typing import Any, Mapping, MutableMapping, Sequence

import openpyxl

from .common import (
    PicklyPyConfigError, PicklyPyVolumeError, as_number, as_str,
    format_csv_rows, integer_chop, make_inventory_row,
)


def split_dilution(compound: str):
    """Recognize a terminal ``-dilNUMBER`` and reject malformed dilution tags."""
    compound = as_str(compound).strip()
    if "-dil" not in compound:
        return None
    base, suffix = compound.rsplit("-dil", 1)
    if not base or "-dil" in base:
        raise PicklyPyConfigError(
            f"Intermediate dilution compound {compound!r} has an empty base or chained dilution tags."
        )
    try:
        factor = float(suffix)
    except (TypeError, ValueError):
        raise PicklyPyConfigError(
            f"Intermediate dilution compound {compound!r} requires a numeric -dil factor."
        ) from None
    if not math.isfinite(factor) or factor <= 1:
        raise PicklyPyConfigError(
            f"Intermediate dilution factor in {compound!r} must be finite and greater than 1."
        )
    return base, factor


def _number(value: Any, context: str, *, allow_zero: bool = False) -> float:
    number = as_number(value)
    if isinstance(value, bool) or number is None or not math.isfinite(number):
        raise PicklyPyConfigError(f"{context} must be a finite number.")
    if number < 0 or (number == 0 and not allow_zero):
        raise PicklyPyConfigError(f"{context} must be {'nonnegative' if allow_zero else 'positive'}.")
    return float(number)


def _file_name(value: Any, suffix: str, context: str) -> str:
    name = as_str(value).strip()
    if not name or name in (".", "..") or name.endswith((".", " ")):
        raise PicklyPyConfigError(f"{context} must contain a filename, without a directory.")
    if re.search(r'[<>:"/\\|?*\x00-\x1f]', name):
        raise PicklyPyConfigError(f"{context} must be a filename without a path or reserved characters.")
    if name.split(".", 1)[0].casefold() in {
        "con", "prn", "aux", "nul", *(f"com{i}" for i in range(1, 10)),
        *(f"lpt{i}" for i in range(1, 10)),
    }:
        raise PicklyPyConfigError(f"{context} cannot use a reserved Windows device filename.")
    if Path(name).suffix:
        if not name.casefold().endswith(suffix.casefold()):
            raise PicklyPyConfigError(f"{context} must have extension {suffix}, or no extension.")
    else:
        name += suffix
    if len(name) > 255:
        raise PicklyPyConfigError(f"{context} filename is too long.")
    return name


def _unique(values):
    return list(dict.fromkeys(as_str(v) for v in values))


def _group_sources(rows):
    groups: dict[str, list] = OrderedDict()
    for row in rows:
        groups.setdefault(as_str(row[0]), []).append(list(row))
    return [row for group in groups.values() for row in group]


@dataclass(frozen=True)
class IntermediateAllocation:
    well: str
    stock: float
    volume_ul: float
    lib_index: int
    identifier: str
    source_plate: str
    source_well: str
    source_wells: tuple[str, ...]
    source_stock: float
    factor: float
    achieved_factor: float
    preload_ul: float
    preparation_nl: float
    preparation_chunks_nl: tuple[float, ...]


class IntermediateDilutions:
    def __init__(self, xlsx_path, default_min_volume_ul):
        self.xlsx_path = Path(xlsx_path).resolve()
        self.allocations: list[IntermediateAllocation] = []
        self._cache: dict[tuple, IntermediateAllocation] = {}
        self._source_definitions: dict[tuple[str, str], tuple[str, float]] = {}
        self._lib_indices: dict[str, list[int]] = {}
        self._source_volume_before_preparation: dict[tuple[str, str], float] = {}
        self.defaults: dict[str, float] = {}
        workbook = openpyxl.load_workbook(self.xlsx_path, data_only=True)
        try:
            candidates = [s for s in workbook if s.title.casefold() == "intermediate_dilutions"]
            if len(candidates) != 1:
                raise PicklyPyConfigError(
                    "--intermediate_dilution requires one Intermediate_dilutions worksheet."
                )
            sheet = candidates[0]
            self.sheet_name = sheet.title
            fields: dict[str, tuple[int, Any]] = {}
            for row in sheet.iter_rows():
                label = as_str(row[0].value).strip().casefold()
                if not label:
                    continue
                if label in fields and label in {
                    "inventory_dil:", "inventory_dil_rack#:", "process_dil:",
                    "picklist name:", "barcode_int:", "plate map:",
                    "minimum volume (ul):", "maximum volume (ul):", "maximum transfer (nl):",
                }:
                    raise PicklyPyConfigError(f"Duplicate Intermediate_dilutions field {label!r}.")
                fields[label] = (row[0].row, row[1].value if len(row) > 1 else None)

            def required(label):
                item = fields.get(label.casefold())
                if item is None or as_str(item[1]).strip() == "":
                    raise PicklyPyConfigError(f"Intermediate_dilutions is missing required field {label}")
                return item[1]

            self.inventory_name = _file_name(required("Inventory_DIL:"), ".csv", "Inventory_DIL:")
            self.process_name = _file_name(required("Process_DIL:"), ".txt", "Process_DIL:")
            self.picklist_name = _file_name(required("Picklist Name:"), ".csv", "Intermediate Picklist Name:")
            rack = _number(required("Inventory_DIL_Rack#:"), "Inventory_DIL_Rack#:")
            if rack != int(rack):
                raise PicklyPyConfigError("Inventory_DIL_Rack#: must be a positive integer.")
            self.inventory_rack = int(rack)
            barcodes = _unique(x.strip() for x in re.split(r"[,;\r\n]+", as_str(required("Barcode_INT:"))) if x.strip())
            if len(barcodes) != 1:
                raise PicklyPyConfigError("Barcode_INT: must identify exactly one intermediate plate barcode.")
            self.barcode = barcodes[0]

            def optional(label, default, allow_zero=False):
                item = fields.get(label.casefold())
                value = item[1] if item is not None else None
                if value is None or as_str(value).strip() == "":
                    self.defaults[label] = float(default)
                    value = default
                return _number(value, label, allow_zero=allow_zero)

            self.minimum_volume_ul = optional("Minimum volume (ul):", default_min_volume_ul, True)
            self.maximum_volume_ul = optional("Maximum volume (ul):", 12.0)
            self.maximum_transfer_nl = optional("Maximum transfer (nl):", 500.0)
            if self.maximum_volume_ul <= self.minimum_volume_ul:
                raise PicklyPyConfigError("Intermediate maximum volume must exceed its minimum volume.")
            if self.maximum_transfer_nl < 2.5:
                raise PicklyPyConfigError("Intermediate maximum transfer must permit at least one 2.5 nL droplet.")
            if "plate map:" not in fields:
                raise PicklyPyConfigError("Intermediate_dilutions is missing its Plate Map: section.")
            map_row = fields["plate map:"][0]
            self._map_row = map_row
            self.eligible_wells: list[tuple[str, float]] = []
            for row in range(16):
                for col in range(24):
                    value = sheet.cell(map_row + 1 + row, 2 + col).value
                    if value is None or isinstance(value, bool):
                        continue
                    number = as_number(value)
                    if number is None:
                        continue  # Text annotations are not eligible preload wells.
                    well = f"{chr(65 + row)}{col + 1}"
                    if not math.isfinite(number) or number < 0:
                        raise PicklyPyConfigError(f"Intermediate preload at {well} must be finite and nonnegative.")
                    if number > 0:
                        self.eligible_wells.append((well, float(number)))
            if not self.eligible_wells:
                raise PicklyPyConfigError("Intermediate Plate Map: has no positive numeric preload wells.")
        finally:
            workbook.close()

    def allocate(self, source_plate, wells, stock, lib_index, identifier, factor,
                 volumes_ul: MutableMapping[tuple[str, str], float], source_minimum_ul):
        source_plate = as_str(source_plate)
        source_wells = tuple(dict.fromkeys(as_str(w) for w in wells))
        identifier = as_str(identifier)
        stock = _number(stock, "Original source stock concentration")
        factor = _number(factor, "Intermediate dilution factor")
        if factor <= 1:
            raise PicklyPyConfigError("Intermediate dilution factor must exceed 1.")
        minimum = _number(source_minimum_ul, "Original source minimum volume", allow_zero=True)
        if source_plate == self.barcode:
            raise PicklyPyConfigError("Barcode_INT: must differ from every original source plate barcode.")
        if not source_plate or not source_wells or any(not w for w in source_wells):
            raise PicklyPyConfigError("Intermediate dilution requires a resolved original source plate and well.")
        # Repeated LIB rows are destination requests, not new physical stock.
        # Reject conflicting descriptions even when rows occur on later plates.
        for source_well in source_wells:
            physical_source = (source_plate, source_well)
            definition = (identifier, stock)
            previous = self._source_definitions.get(physical_source)
            if previous is not None and previous != definition:
                raise PicklyPyConfigError(
                    f"Conflicting repeated stock at {source_plate}:{source_well}: "
                    f"compound/stock {previous!r} versus {definition!r}. "
                    "Repeated LIB entries for one physical well must agree."
                )
            self._source_definitions[physical_source] = definition
        # Look up the actual preparation source well, not its LIB row number,
        # placeholder, or an ordered list of interchangeable source candidates.
        # Reuse precedes volume/capacity checks: no second preparation is needed.
        for source_well in source_wells:
            key = (source_plate, source_well, identifier, stock, factor)
            if key in self._cache:
                return self._record_use(self._cache[key], lib_index)
        if len(self.allocations) >= len(self.eligible_wells):
            raise PicklyPyConfigError(
                f"Intermediate plate {self.barcode} has insufficient eligible preload wells: "
                f"capacity {len(self.eligible_wells)}; another compound/factor requires a well."
            )
        well, preload = self.eligible_wells[len(self.allocations)]
        units = (Decimal(str(preload)) * 1000 / (Decimal(str(factor)) - 1) / Decimal("2.5"))
        prep_nl = float(units.quantize(Decimal("1"), rounding=ROUND_HALF_UP) * Decimal("2.5"))
        if prep_nl <= 0:
            raise PicklyPyVolumeError(f"Intermediate preparation for {identifier} rounds to 0 nL at factor {factor:g}.")
        prepared = preload + prep_nl / 1000.0
        if prepared > self.maximum_volume_ul + 1e-12 or prepared <= self.minimum_volume_ul:
            raise PicklyPyVolumeError(
                f"Intermediate {self.barcode}:{well} prepared volume {prepared:g} uL must be "
                f"above minimum {self.minimum_volume_ul:g} and at most {self.maximum_volume_ul:g} uL."
            )
        chosen = next((w for w in source_wells if volumes_ul.get((source_plate, w), 0.0)
                       + 1e-12 >= minimum + prep_nl / 1000.0), None)
        if chosen is None:
            raise PicklyPyVolumeError(f"Out of Volume preparing intermediate dilution for {identifier} from {source_plate}.")
        int_key = (self.barcode, well)
        if int_key in volumes_ul:
            raise PicklyPyConfigError(f"Intermediate allocation would overwrite existing source volume at {self.barcode}:{well}.")
        chunk_units = int(math.floor(self.maximum_transfer_nl / 2.5))
        remaining_units = int(round(prep_nl / 2.5))
        chunks = []
        while remaining_units:
            take = min(chunk_units, remaining_units)
            chunks.append(take * 2.5)
            remaining_units -= take
        actual_stock = stock * (prep_nl / 1000.0) / prepared
        allocation = IntermediateAllocation(
            well=well, stock=actual_stock, volume_ul=prepared, lib_index=int(lib_index),
            identifier=identifier, source_plate=source_plate, source_well=chosen,
            source_wells=source_wells, source_stock=stock, factor=factor,
            achieved_factor=prepared / (prep_nl / 1000.0), preload_ul=preload,
            preparation_nl=prep_nl, preparation_chunks_nl=tuple(chunks),
        )
        self._source_volume_before_preparation.setdefault(
            (source_plate, chosen), volumes_ul[(source_plate, chosen)]
        )
        volumes_ul[(source_plate, chosen)] -= prep_nl / 1000.0
        volumes_ul[int_key] = prepared
        self.allocations.append(allocation)
        key = (source_plate, chosen, identifier, stock, factor)
        self._cache[key] = allocation
        return self._record_use(allocation, lib_index)

    def _record_use(self, allocation, lib_index):
        """Keep every requesting LIB row while exporting one physical preparation."""
        indices = self._lib_indices.setdefault(allocation.well, [])
        if int(lib_index) not in indices:
            indices.append(int(lib_index))
        return allocation

    def _lib_rows(self, lib, references, volumes_ul):
        width = max((len(row) for row in lib), default=0)
        headers = list(lib[0]) if lib else []
        headers += [""] * (width - len(headers))
        extra_headers = [
            "Intermediate well", "Intermediate plate barcode", "Intermediate stock concentration",
            "Intermediate prepared volume (uL)", "Intermediate remaining volume (uL)",
            "Preloaded diluent (uL)", "Requested dilution factor", "Achieved dilution factor",
            "Preparation transfer (nL)", "Original source well", "Original source plate barcode",
            "Original stock concentration", "Original LIB row", "Original LIB worksheet rows",
        ]
        rows = [headers + extra_headers]
        for allocation in self.allocations:
            if not 0 < allocation.lib_index < len(lib):
                raise PicklyPyConfigError(f"Invalid original LIB row {allocation.lib_index} for intermediate export.")
            row = list(lib[allocation.lib_index]) + [""] * (width - len(lib[allocation.lib_index]))
            for role, value in (("well", allocation.well), ("barcode", self.barcode), ("stock", allocation.stock)):
                for col in getattr(references, role, ()):
                    if col < 0 or col >= width:
                        raise PicklyPyConfigError(f"Intermediate LIB {role} reference is outside the LIB table.")
                    row[col] = integer_chop(value) if isinstance(value, float) else value
            row.extend([
                allocation.well, self.barcode, allocation.stock, allocation.volume_ul,
                volumes_ul.get((self.barcode, allocation.well), allocation.volume_ul),
                allocation.preload_ul, allocation.factor, allocation.achieved_factor,
                allocation.preparation_nl, allocation.source_well, allocation.source_plate,
                allocation.source_stock, allocation.lib_index,
                ", ".join(str(i + 1) for i in self._lib_indices[allocation.well]),
            ])
            rows.append(row)
        return rows

    def write_outputs(self, out_dir, header, cell_rows, src_unique, dst_unique, lib, references, volumes_ul):
        """Plan all outputs before writing; preserve preparation/direct/diluted phases."""
        out_dir = Path(out_dir).resolve()
        if not out_dir.is_dir():
            raise PicklyPyConfigError(f"Intermediate output directory does not exist: {out_dir}")
        dst_unique = _unique(dst_unique)
        if self.barcode in dst_unique:
            raise PicklyPyConfigError("Barcode_INT: must differ from every final destination barcode.")
        original_sources = _unique([a.source_plate for a in self.allocations]
                                   + [s for s in src_unique if as_str(s) != self.barcode])
        prep_rows = []
        for allocation in self.allocations:
            row = [allocation.source_plate, allocation.source_well, self.barcode]
            for chunk in allocation.preparation_chunks_nl:
                row.extend([integer_chop(chunk), allocation.well])
            prep_rows.append(row)
        prep_rows = _group_sources(prep_rows)
        direct_rows = _group_sources([r for r in cell_rows if as_str(r[0]) != self.barcode])
        diluted_rows = _group_sources([r for r in cell_rows if as_str(r[0]) == self.barcode])
        ordered_cell_rows = direct_rows + diluted_rows
        planned: dict[str, dict] = OrderedDict()
        reserved = {
            "table merge.txt", "plate map compounds.txt", "plate map concentrations.txt",
            f"{self.xlsx_path.stem}_reformatted_LIB.xlsx".casefold(),
        }

        def add(name, kind, rows):
            suffix = ".txt" if kind == "process" else ".csv"
            name = _file_name(name, suffix, f"{kind} output")
            folded = name.casefold()
            if folded in reserved:
                raise PicklyPyConfigError(f"Intermediate output {name!r} conflicts with Screen metadata.")
            path = out_dir / name
            if path.resolve().parent != out_dir or path.resolve() == self.xlsx_path:
                raise PicklyPyConfigError(f"Output filename {name!r} escapes its directory or overwrites the input design.")
            old = planned.get(folded)
            if old is not None and old["kind"] != kind:
                raise PicklyPyConfigError(f"Output {name!r} is used for incompatible {old['kind']} and {kind} files.")
            target = planned.setdefault(folded, {"path": path, "kind": kind, "rows": []})
            target["rows"].extend(rows)

        # Insert prep first: if names match, this order is preserved in one file.
        add(self.picklist_name, "picklist", prep_rows)
        add(header.picklist_name, "picklist", ordered_cell_rows)
        add(header.inventory_src_name, "inventory", [make_inventory_row(b, header.inventory_src_rack, i + 1)
                                                       for i, b in enumerate(original_sources)])
        add(header.inventory_dst_name, "inventory", [make_inventory_row(b, header.inventory_dst_rack, i + 1)
                                                       for i, b in enumerate(dst_unique)])
        add(self.inventory_name, "inventory", [make_inventory_row(self.barcode, self.inventory_rack, 1)])
        add(header.process_src_name, "process", original_sources)
        add(header.process_dst_name, "process", dst_unique)
        add(self.process_name, "process", [self.barcode])
        payloads: dict[Path, bytes] = OrderedDict()
        for target in planned.values():
            rows = target["rows"]
            if target["kind"] == "inventory":
                by_barcode = {}
                by_address = {}
                for row in rows:
                    barcode = as_str(row[0])
                    address = tuple(row[7:9])
                    if barcode in by_barcode and by_barcode[barcode] != row:
                        raise PicklyPyConfigError(f"Conflicting inventory addresses for {barcode} in {target['path'].name}.")
                    if address in by_address and by_address[address] != barcode:
                        raise PicklyPyConfigError(f"Conflicting inventory rack/position {address} in {target['path'].name}.")
                    by_barcode[barcode] = row
                    by_address[address] = barcode
                text = format_csv_rows(list(by_barcode.values()), quote=True)
            elif target["kind"] == "process":
                text = "\n".join(_unique(rows))
            else:
                text = format_csv_rows(rows, quote=True)
            payloads[target["path"]] = text.encode("utf-8")

        lib_rows = self._lib_rows(lib, references, volumes_ul)
        lib_book = openpyxl.Workbook()
        lib_sheet = lib_book.active
        lib_sheet.title = "LIB"
        for row in lib_rows:
            lib_sheet.append(row)
        lib_buffer = BytesIO()
        lib_book.save(lib_buffer)
        lib_book.close()

        # Screen reads cached formula results. Materializing them preserves a
        # rerunnable generated workbook; saving formulas via openpyxl loses caches.
        design = openpyxl.load_workbook(self.xlsx_path, data_only=True)
        try:
            for name in ("Intermediate_LIB", "Intermediate_Allocation"):
                if name in design.sheetnames:
                    del design[name]
            library_sheet = design.create_sheet("Intermediate_LIB")
            for row in lib_rows:
                library_sheet.append(row)
            allocation_sheet = design.create_sheet("Intermediate_Allocation")
            allocation_sheet.append(["Generated intermediate dilution plan"])
            allocation_sheet.append(["Rerun", "Use --intermediate_dilution true; allocation is rebuilt from the original LIB and preload map."])
            allocation_sheet.append(["Formula handling", "Cached formula results are materialized as values; the original input workbook is unchanged."])
            allocation_sheet.append(["Source design", self.xlsx_path.name])
            allocation_sheet.append(["Intermediate layout", "Fixed row-major eligible preload wells; final DST randomization does not move these wells."])
            allocation_sheet.append([])
            for row in lib_rows:
                allocation_sheet.append(row)
            intermediate_sheet = design[self.sheet_name]
            next_default_row = max(intermediate_sheet.max_row + 1, self._map_row + 17)
            for label, value in self.defaults.items():
                matches = [r for r in range(1, intermediate_sheet.max_row + 1)
                           if as_str(intermediate_sheet.cell(r, 1).value).strip().casefold() == label.casefold()]
                if matches:
                    intermediate_sheet.cell(matches[0], 2, value)
                else:
                    # A sparse preload map may end at A1 in the XLSX. Appending
                    # immediately after max_row would turn defaults in column B
                    # into unintended new numeric preload wells on the next run.
                    intermediate_sheet.cell(next_default_row, 1, label)
                    intermediate_sheet.cell(next_default_row, 2, value)
                    next_default_row += 1
            design_buffer = BytesIO()
            design.save(design_buffer)
        finally:
            design.close()

        preparation_debits: dict[tuple[str, str], float] = {}
        for allocation in self.allocations:
            key = (allocation.source_plate, allocation.source_well)
            preparation_debits[key] = preparation_debits.get(key, 0.0) + allocation.preparation_nl
        source_balances = [
            {
                "source_plate": plate, "source_well": well,
                "remaining_volume_ul": remaining,
                "preparation_transfer_nl": preparation_debits.get((plate, well), 0.0),
                "volume_before_first_preparation_ul": self._source_volume_before_preparation.get((plate, well)),
            }
            for (plate, well), remaining in volumes_ul.items() if plate != self.barcode
        ]
        report = {
            "input": str(self.xlsx_path), "intermediate_barcode": self.barcode,
            "eligible_well_count": len(self.eligible_wells), "allocated_well_count": len(self.allocations),
            "source_barcodes": original_sources, "destination_barcodes": dst_unique,
            "minimum_volume_ul": self.minimum_volume_ul, "maximum_volume_ul": self.maximum_volume_ul,
            "maximum_transfer_nl": self.maximum_transfer_nl, "recorded_defaults": self.defaults,
            "formula_handling": "Generated design materializes cached values; original workbook unchanged.",
            "phase_order": ["Preparation", "Original sources to cells", "Intermediate to cells"],
            "original_source_balances": source_balances,
            "allocations": [{**asdict(a),
                             "lib_indices": self._lib_indices[a.well],
                             "lib_worksheet_rows": [i + 1 for i in self._lib_indices[a.well]],
                             "remaining_volume_ul": volumes_ul.get((self.barcode, a.well), a.volume_ul)}
                            for a in self.allocations],
        }
        extras = {
            f"{self.xlsx_path.stem}_intermediate_LIB.xlsx": lib_buffer.getvalue(),
            f"{self.xlsx_path.stem}_intermediate_design.xlsx": design_buffer.getvalue(),
            f"{self.xlsx_path.stem}_intermediate_report.json": (json.dumps(report, indent=2) + "\n").encode("utf-8"),
        }
        for name, content in extras.items():
            path = out_dir / name
            if path.resolve().parent != out_dir or path.resolve() == self.xlsx_path or name.casefold() in planned:
                raise PicklyPyConfigError(f"Generated intermediate output {name!r} conflicts with another output or the input.")
            payloads[path] = content
        for path, content in payloads.items():
            path.write_bytes(content)
            print(f"Saving: {path}")
        report["outputs"] = [str(path) for path in payloads]
        return report
